"""生成 E2E / 单测使用的 Word 夹具（一次性运行，产物已提交）。

运行：$MIMO_PYTHON tests/fixtures/make-sample-docx.py
"""
from pathlib import Path

from docx import Document

OUT = Path(__file__).with_name("sample.docx")

doc = Document()
doc.add_heading("MiMo Attachment Fixture", level=1)
doc.add_paragraph("中文文档测试：这是一段用于端到端测试的正文。")
doc.add_paragraph("Second paragraph with English text.")
doc.save(str(OUT))
print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
